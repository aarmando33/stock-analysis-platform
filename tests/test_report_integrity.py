import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import price_feed
from buy_sell_monitor import calculate, run
from export_monitor import build
import test_buy_sell_monitor as fixtures
history=fixtures.history


class FeedIntegrity(unittest.TestCase):
    def test_crypto_daily_bar_must_close_at_utc_midnight(self):
        g=self.frame();g.Symbol='BTC-USD'
        before=price_feed.audit_history(g,['BTC-USD'],'2026-10-02',now='2026-10-02T22:00:00Z')
        after=price_feed.audit_history(g,['BTC-USD'],'2026-10-02',now='2026-10-03T00:00:00Z')
        self.assertEqual(before.price_status.iloc[0],'UNCOMPLETED_CRYPTO_CANDLE')
        self.assertEqual(after.price_status.iloc[0],'OK')
    def test_full_universe_feed_to_workbook_integration(self):
        symbols=price_feed.tickers()
        self.assertEqual(len(symbols),168)
        frames=[]
        for symbol in symbols:
            if symbol=='CURLD':continue
            g=history(n=30);g['Provider']='fixture';g.Symbol=symbol;frames.append(g)
        with tempfile.TemporaryDirectory() as directory:
            output=Path(directory)
            with patch('price_feed.OUT',output), patch('price_feed.exclusive_end',return_value=pd.Timestamp('2026-10-03').to_pydatetime()), patch('price_feed.yahoo',return_value=pd.concat(frames)), patch('price_feed.tiingo',return_value=price_feed.empty()):
                price_feed.main()
            report=run(output,price_feed.ROOT/'tickers.csv','2026-10-02',historical=True)
            self.assertEqual(report['counts']['Expected'],168)
            self.assertEqual(report['counts']['Usable'],167)
            self.assertEqual(len(report['views']['Master']),168)
            self.assertEqual(report['views']['Master'][-1]['price_status'] in ['OK','MISSING_PRICE'],True)
            workbook=build(report);path=output/'monitor.xlsx';workbook.save(path)
            saved=load_workbook(path)
            self.assertEqual(saved['Master'].max_row,169)
            summary={r[0].value:r[1].value for r in saved['Summary'].iter_rows(min_row=2)}
            self.assertEqual(summary['Usable'],167)

    def test_alias_preserves_identity_and_effective_date(self):
        self.assertEqual(price_feed.provider_symbols(['CURLD'],pd.Timestamp('2026-10-03')),{'CURLD':'CURLF'})
        self.assertEqual(price_feed.provider_symbols(['CURLD'],pd.Timestamp('2026-06-10')),{'CURLD':'CURLD'})
    def frame(self):
        g=history();g['Provider']='fixture'
        return g

    def test_all_stale_feed_is_not_ok(self):
        x=price_feed.audit_history(self.frame(),['TEST'],'2026-10-05')
        self.assertEqual(x.price_status.iloc[0],'STALE_PRICE')
        self.assertTrue(pd.isna(x['Primary Win%'].iloc[0]))

    def test_calendar_early_close_and_holiday(self):
        self.assertEqual(price_feed.exclusive_end('2026-11-27T18:30:00Z').date().isoformat(),'2026-11-28')
        self.assertEqual(price_feed.exclusive_end('2026-07-03T22:00:00Z').date().isoformat(),'2026-07-03')

    def test_duplicate_and_invalid_ohlc_rejected(self):
        g=self.frame()
        for col,value in [('Open',float('inf')),('Open',10000),('Close',0)]:
            with self.subTest(col=col,value=value):
                bad=g.copy();bad.loc[0,col]=value
                self.assertEqual(price_feed.audit_history(bad,['TEST'],'2026-10-02').price_status.iloc[0],'INVALID_HISTORY')
        bad=pd.concat([g,g.iloc[[0]]])
        self.assertEqual(price_feed.audit_history(bad,['TEST'],'2026-10-02').price_status.iloc[0],'INVALID_HISTORY')

    def test_tiingo_selects_adjusted_fields_without_duplicate_columns(self):
        raw={'date':'2026-10-02','open':200,'high':220,'low':180,'close':210,'volume':100,
             'adjOpen':100,'adjHigh':110,'adjLow':90,'adjClose':105,'adjVolume':200}
        response=unittest.mock.Mock();response.json.return_value=[raw]
        with patch.dict('os.environ',{'TIINGO_API_KEY':'fixture'}), patch('price_feed.requests.get',return_value=response), patch('price_feed.time.sleep'):
            result=price_feed.tiingo(['TEST'],'2020-03-01',pd.Timestamp('2026-10-03').to_pydatetime())
        self.assertFalse(result.columns.duplicated().any())
        self.assertEqual(result.Close.iloc[0],105)

    def test_rsi_monotonic_up_has_value(self):
        self.assertEqual(price_feed.wilder_rsi(pd.Series(np.arange(30.))).iloc[-1],100)


class AdditionalPipeline(unittest.TestCase):
    setUp=fixtures.Pipeline.setUp
    tearDown=fixtures.Pipeline.tearDown
    go=fixtures.Pipeline.go

    def test_nonfinite_or_zero_volume_does_not_confirm_flow(self):
        for value in [float('inf'),0.]:
            g=history();g.Volume=value
            result=calculate(g,g.Close.iloc[-1],'2026-10-02')
            self.assertEqual(result['Money Flow'],'Unavailable')
            self.assertEqual(result['Volume Confirmation'],'Unavailable')

    def test_interior_missing_close_is_not_silently_removed(self):
        g=pd.read_csv(self.p/'price_history.csv.gz');g.loc[g.Date.eq('2026-10-01'),'Close']=np.nan
        g.to_csv(self.p/'price_history.csv.gz',index=False,compression='gzip')
        self.assertEqual(self.go()['views']['Master'][0]['price_status'],'INCOMPLETE_HISTORY')

    def test_equity_weekend_padding_does_not_create_false_gap(self):
        g=pd.read_csv(self.p/'price_history.csv.gz')
        padding=g.iloc[[0]].copy();padding['Date']='2026-09-26'
        padding[['Open','High','Low','Close','Volume']]=np.nan
        pd.concat([g,padding]).to_csv(self.p/'price_history.csv.gz',index=False,compression='gzip')
        self.assertEqual(self.go()['counts']['Usable'],1)
    def test_audit_read_before_missing_price_files(self):
        (self.p/'prices_latest.csv').unlink()
        (self.p/'price_audit.json').write_text(json.dumps({'latest_market_date':'2026-10-01'}))
        self.assertEqual(self.go()['failure']['code'],'FEED_SESSION_MISMATCH')

    def test_missing_audit_structured(self):
        (self.p/'price_audit.json').unlink()
        self.assertEqual(self.go()['failure']['code'],'INVALID_OR_MISSING_AUDIT')

    def test_infinite_and_outside_open_rejected(self):
        g=pd.read_csv(self.p/'price_history.csv.gz')
        for value in [float('inf'),10000]:
            bad=g.copy();bad.loc[0,'Open']=value
            bad.to_csv(self.p/'price_history.csv.gz',index=False,compression='gzip')
            self.assertEqual(self.go()['counts']['Usable'],0)

    def test_blank_context_date_rejected(self):
        f=self.p/'context.csv'
        pd.DataFrame([{'Ticker':'TEST','Source':'fixture','As Of':None,'Raw Evidence':'fixture'}]).to_csv(f,index=False)
        with self.assertRaisesRegex(ValueError,'cannot be blank'):self.go(context_path=f)

    def test_failed_price_keeps_owned_position(self):
        f=self.p/'positions.csv'
        pd.DataFrame([{'Ticker':'TEST','Source':'fixture','As Of':'2026-10-02','Confirmed':True,'Qty':10,'Cost Basis':100}]).to_csv(f,index=False)
        lat=pd.read_csv(self.p/'prices_latest.csv');lat.price_status='MISSING_PRICE';lat.to_csv(self.p/'prices_latest.csv',index=False)
        x=self.go(positions_path=f)
        self.assertEqual(len(x['views']['Owned Positions']),1)
        self.assertEqual(x['views']['Owned Positions'][0]['Overall Signal/Action'],'Data Unavailable')
        self.assertEqual(x['views']['Owned Positions'][0]['Qty'],10)
        self.assertNotIn('Current Value',x['views']['Owned Positions'][0])

    def test_xlsx_round_trip_matches_every_calculation(self):
        report=self.go()
        path=self.p/'monitor.xlsx';build(report).save(path)
        wb=load_workbook(path,data_only=False)
        headers=[c.value for c in wb['Calculations'][1]]
        for cells,expected in zip(wb['Calculations'].iter_rows(min_row=2),report['calculations']):
            for key,cell in zip(headers,cells):
                value=expected[key]
                if isinstance(value,(float,int)) and not isinstance(value,bool):
                    self.assertAlmostEqual(cell.value,value)
                else:self.assertEqual(cell.value,value)
        self.assertEqual(len(wb.sheetnames),len(report['views'])+4)
        self.assertEqual(wb['Master'].max_row,2)
        self.assertEqual(wb['Master'].freeze_panes,'C2')
        self.assertIn('code_sha256',report)
        self.assertIn('universe',report['input_sha256'])

    def test_failure_workbook_has_reason_and_no_rankings(self):
        (self.p/'price_audit.json').write_text(json.dumps({'latest_market_date':'2026-10-01'}))
        report=self.go();wb=build(report)
        summary={r[0].value:r[1].value for r in wb['Summary'].iter_rows(min_row=2)}
        self.assertEqual(summary['Failure code'],'FEED_SESSION_MISMATCH')
        self.assertEqual(wb['Top Opportunities']['A2'].value,'No qualifying records')

    def test_source_strings_never_become_formulas(self):
        report=self.go();report['calculations'][0]['Context Source']='=1+1'
        wb=build(report)
        sheet=wb['Calculations'];headers=[c.value for c in sheet[1]]
        cell=sheet.cell(2,headers.index('Context Source')+1)
        self.assertEqual(cell.data_type,'s')

    def test_direct_calculator_ignores_future_benchmarks(self):
        g=history();bm=g.set_index('Date').Close
        future=pd.Series([100000.],index=[pd.Timestamp('2026-10-05')])
        a=calculate(g,g.Close.iloc[-1],'2026-10-02',benchmark=bm)
        b=calculate(g,g.Close.iloc[-1],'2026-10-02',benchmark=pd.concat([bm,future]))
        self.assertEqual(a['Environment Points'],b['Environment Points'])
        self.assertEqual(a['RS vs Market 20D %'],b['RS vs Market 20D %'])


if __name__=='__main__':unittest.main()
