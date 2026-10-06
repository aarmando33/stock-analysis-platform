"""Versioned Python workbook exporter for Stock Buy/Sell Monitor v2.

Uses a portable local XLSX writer; no background service or socket is required.
All calculations arrive precomputed from buy_sell_monitor.py; this module is presentation-only.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VISIBLE = [
    'Overall Rank','Ticker','Owned/Watch','Price','Price Date','Buy Zone',
    'Primary Win%','52W Closing-Range Win%','Opportunity Score','Setup Confidence',
    'Setup Confidence Coverage %','Evidence Coverage %','Bottom Confidence',
    'Bottom Confidence Coverage %','Research Coverage %','Overall Signal/Action','Action Reason',
    'Short Support','Support Timeframe','Major Support','Distance to Support %',
    'Short Resistance','Major Resistance','Distance to Resistance %',
    'Invalidation','Downside %','Upside %','Risk/Reward','Trim Zone',
    'Recent Volume','Average Volume 20D','Relative Volume','20D Net Volume %',
    'Basing Status','Bottom/Falling-Knife Status','Volume Confirmation','Rotation Stage','Money Flow'
]
OWNED_EXTRA = ['Qty','Cost Basis','Current Value','Unrealized $','Unrealized %']


def col_letter(n: int) -> str:
    s=''
    while n:
        n, rem = divmod(n-1, 26)
        s=chr(65+rem)+s
    return s


def sheet(wb, name, rows, headers=None):
    sh = wb.create_sheet(name)
    headers = headers or list(dict.fromkeys(k for r in rows for k in r)) or ['Status']
    body = [[r.get(h) for h in headers] for r in rows] if rows else [['No qualifying records'] + [None]*(len(headers)-1)]
    for values in [headers, *body]:
        sh.append(values)
    sh.sheet_view.showGridLines = False
    sh.freeze_panes = 'C2' if len(headers) >= 2 else 'B2'
    sh.auto_filter.ref = sh.dimensions
    for row in sh:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            cell.font = Font(size=11)
            # Source text must stay literal, never execute as an Excel formula.
            if isinstance(cell.value, str):
                cell.data_type = 's'
        sh.row_dimensions[row[0].row].height = 110 if name == 'Methodology' else 60
    for cell in sh[1]:
        cell.fill = PatternFill('solid', fgColor='193F60')
        cell.font = Font(size=11, bold=True, color='FFFFFF')
    sh.row_dimensions[1].height = 52
    for i, h in enumerate(headers, 1):
        width = 36 if any(x in h for x in ['Reason','Source','Status','Action','Location','Basis','Evidence','Owned/Watch','Formula','Notes','Timeframe']) else 18
        sh.column_dimensions[get_column_letter(i)].width = width
        fmt = 'General'
        if '%' in h: fmt = '0.0"%"'
        elif any(x in h for x in ['Score','Confidence','Points']) or h in ['Risk/Reward','Relative Volume']: fmt = '0.0'
        elif h.endswith(('Price','Support','Resistance','MA','Basis','Value')) or h in ['Price','ATR','Invalidation','Unrealized $'] or h.startswith('MACD'): fmt = '0.00'
        elif 'Volume' in h and 'Ratio' not in h and 'Confirmation' not in h: fmt = '#,##0'
        for cells in sh.iter_rows(min_row=2, min_col=i, max_col=i):
            cells[0].number_format = fmt
    return sh


def build(report):
    wb=Workbook()
    wb.remove(wb.active)
    failure=report.get('failure') or {}
    summary=[
        {'Metric':'Research coverage','Value':report.get('research_status','Unavailable'),'Notes':'Missing fundamental/sector/benchmark evidence cannot support buy decisions'},
        {'Metric':'Report status','Value':report.get('status'),'Notes':'Historical rebuild; not current trading guidance' if report.get('historical') else 'Latest completed session'},
        {'Metric':'Market session','Value':report.get('session'),'Notes':'All technical history is cut off at this session'},
    ]
    for k,v in report.get('counts',{}).items():
        summary.append({'Metric':k,'Value':v,'Notes':'Computed from validated records'})
    if failure:
        summary += [
            {'Metric':'Failure code','Value':failure.get('code'),'Notes':failure.get('detail')},
            {'Metric':'Requested session','Value':failure.get('requested_session'),'Notes':'Freshness gate request'},
            {'Metric':'Feed session','Value':failure.get('feed_session'),'Notes':'Artifact latest_market_date'},
        ]
    summary += [
        {'Metric':'Ownership','Value':'Confirmed source supplied' if report.get('positions') else 'Unavailable','Notes':'No hardcoded holdings used'},
        {'Metric':'Scoring','Value':'Fixed 20/25/20/20/10/5 weights','Notes':'Missing evidence earns zero points; coverage is reported separately'},
        {'Metric':'52W Win% label','Value':'52W Closing-Range Win%','Notes':'Closing-price range-position metric, not probability of profit'},
        {'Metric':'Volume flow label','Value':'20D Net Volume %','Notes':'20-session signed volume divided by total volume'},
    ]
    sheet(wb,'Summary',summary,['Metric','Value','Notes'])

    for name,rows in report.get('views',{}).items():
        headers=list(VISIBLE)
        if name=='Owned Positions':
            headers+=OWNED_EXTRA
        sheet(wb,name,rows,headers)

    calculations=report.get('calculations',[])
    sheet(wb,'Calculations',calculations)

    audit_cols=[
        'Ticker','Universe','price_status','Price Date','Price Provider','Provider Symbol','Price Basis',
        'Context Source','Context As Of','Ownership Source',
        'Short Support','Short Support Source','Short Support Date','Short Support Strength','Short Support Tests','Support Timeframe','Support Defended',
        'Major Support','Major Support Source','Major Support Date','Major Support Strength','Major Support Tests',
        'Short Resistance','Short Resistance Source','Short Resistance Date','Major Resistance','Major Resistance Source','Major Resistance Date',
        'Context Raw Evidence'
    ]
    audit_rows=[{k:r.get(k) for k in audit_cols} for r in calculations]
    audit=sheet(wb,'Audit-Provenance',audit_rows,audit_cols)
    hashes=[['Input/code file','SHA256'],*[[k,v] for group in ['source_sha256','input_sha256','code_sha256'] for k,v in report.get(group,{}).items()]]
    start=max(5,len(audit_rows)+5)
    for row_index, values in enumerate(hashes, start):
        for col_index, value in enumerate(values, 1):
            audit.cell(row_index, col_index, value)

    methodology=[
        {'Metric':'Primary Win%','Formula':'(maximum adjusted close - current)/(maximum - minimum) x 100','Window':'Closing range from 2020-03-01 or first available session'},
        {'Metric':'52W Closing-Range Win%','Formula':'Same closing-range position formula','Window':'Trailing 365 calendar days; explicitly a closing-price range metric'},
        {'Metric':'price_suggest_80','Formula':'20% maximum + 80% minimum adjusted close','Window':'Primary historical window'},
        {'Metric':'YTD','Formula':'Current / prior-year last available adjusted close - 1','Window':'Unavailable without prior-year reference'},
        {'Metric':'Support selection','Formula':'Short: nearest 5/21-session pivot or MA20. Major: strongest 63/126/252-session + MA50/100/200 confluence cluster; distance breaks ties','Window':'Touch episodes within 1%; clusters within 0.5%'},
        {'Metric':'Defended support','Formula':'Selected short support has >=2 independent test episodes, a recent low within 2%, >=2% rebound from that recent low, and current remains above support','Window':'Recent 10 sessions plus short-level test history'},
        {'Metric':'Basing','Formula':'(21D high - 21D low)/21D high <=10% and |20-session return|<=8%; confirmed adds higher low, defended support, declining ATR and volume contraction','Window':'21 sessions'},
        {'Metric':'MACD','Formula':'EMA12 - EMA26; signal EMA9; histogram line - signal','Window':'Adjusted daily close'},
        {'Metric':'ATR','Formula':'Wilder-style EWM of max(high-low, |high-prev close|, |low-prev close|)','Window':'14 sessions'},
        {'Metric':'20D Net Volume %','Formula':'20-session signed volume / total volume x 100; cumulative OBV retained separately','Window':'20 sessions'},
        {'Metric':'Breakout','Formula':'Above prior 21-session high; confirmation requires consecutive prior-session break plus relative volume >=1.5','Window':'Latest session excluded from reference level'},
        {'Metric':'Opportunity Score','Formula':'Fixed 20/25/20/20/10/5 buckets; unavailable evidence earns zero and weights are never renormalized','Window':'Evidence Coverage % reports available weighted evidence'},
        {'Metric':'Confidence','Formula':'Numeric independent evidence score; coverage reported separately','Window':'Setup Confidence and Bottom Confidence each 0-100'},
        {'Metric':'Bottom status','Formula':'Insufficient evidence when research coverage <50%; otherwise evidence bands apply','Window':'Research evidence remains explicitly unavailable when not sourced'},
        {'Metric':'Action priority','Formula':'Breakdown -> owned Trim/Raise Protection -> coverage gate -> breakout/buy -> proximity actions','Window':'Owned risk warnings are not suppressed by missing buy research'},
        {'Metric':'Ownership','Formula':'Only confirmed sourced positions; otherwise quantity/basis/P&L unavailable','Window':'No historical hardcoded positions'},
        {'Metric':'Scope','Formula':'Discovery continuity, anchored VWAP/volume profile, analyst/insider/institutional ingestion and expanded market regime remain outstanding','Window':'Not a claim of full master-spec implementation'},
    ]
    sheet(wb,'Methodology',methodology,['Metric','Formula','Window'])
    return wb


def main():
    p=argparse.ArgumentParser()
    p.add_argument('input')
    p.add_argument('output')
    a=p.parse_args()
    report=json.loads(Path(a.input).read_text(encoding='utf-8'))
    wb=build(report)
    Path(a.output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.output)
    print(a.output)


if __name__=='__main__':
    main()
